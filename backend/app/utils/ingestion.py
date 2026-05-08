import os
import re
import numpy as np
import pandas as pd
import datetime
from io import BytesIO
from dateutil.parser import parse
from openpyxl import load_workbook
from typing import List, Dict, Any, Optional
import logging

logger = logging.getLogger(__name__)

class IngestionEngine:
    """
    Advanced Logic Engine for "Effective Ingestion" of Excel/CSV files.
    Optimized for industrial trackers with merged headers and sparse categorical columns.
    """

    def __init__(self, max_header_scan: int = 8):
        self.max_header_scan = max_header_scan

    def ingest(self, file_content: bytes, file_name: str, _debug: bool = False) -> List[Dict[str, Any]]:
        """Main entry point for ingestion from raw file bytes."""
        
        # Create a file-like object and attach a filename for extension checking
        file_obj = BytesIO(file_content)
        setattr(file_obj, "filename", file_name)

        debug_info = {}

        # ── 1. Read ──────────────────────────────────────────
        df = self._read_file(file_obj)
        debug_info["step1_shape_after_read"] = list(df.shape)

        # ── 2. Trim ──────────────────────────────────────────
        df = self._trim_dataframe(df)
        debug_info["step2_shape_after_trim"] = list(df.shape)

        if df.empty:
            debug_info["exit"] = "empty after trim"
            return ([], debug_info) if _debug else []

        # ── 3. Detect header rows ────────────────────────────
        header_rows = self._detect_header_rows(df)
        debug_info["step3_header_rows"] = header_rows

        # ── 4. Expand merged header cells ────────────────────
        # Seek back so openpyxl can re-read merge ranges
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        df = self._expand_header_merged_cells(df, header_rows, file_obj=file_obj)

        # ── 5. Build column names ────────────────────────────
        headers = self._build_headers(df, header_rows)
        debug_info["step5_headers"] = headers
        debug_info["step5_header_count"] = len(headers)
        debug_info["step5_df_col_count"] = df.shape[1]

        # ── 6. Detect data start ─────────────────────────────
        data_start = self._detect_data_start(df, header_rows)
        debug_info["step6_data_start_row"] = data_start
        debug_info["step6_total_rows"] = len(df)

        if data_start >= len(df):
            debug_info["exit"] = "data_start beyond end of df"
            return ([], debug_info) if _debug else []

        # ── 7. Slice data + assign headers ───────────────────
        data_df = df.iloc[data_start:].reset_index(drop=True).copy()
        debug_info["step7_data_rows"] = len(data_df)

        if data_df.empty:
            debug_info["exit"] = "data slice empty"
            return ([], debug_info) if _debug else []

        # Guard: column count mismatch
        n_df_cols = data_df.shape[1]
        n_headers = len(headers)

        if n_df_cols != n_headers:
            if n_df_cols < n_headers:
                headers = headers[:n_df_cols]
            else:
                for extra in range(n_headers, n_df_cols):
                    headers.append(f"column_{extra}")

            debug_info["step7_col_mismatch_reconciled"] = {
                "df_cols": n_df_cols,
                "header_count": n_headers,
                "final_headers": headers,
            }

        data_df.columns = headers

        # ── 8. Forward-fill merged body cells ────────────────
        data_df = self._fill_merged_data(data_df)

        # ── 9. Drop all-empty rows ───────────────────────────
        data_df = data_df.dropna(how="all")
        debug_info["step9_rows_after_dropna"] = len(data_df)

        if data_df.empty:
            debug_info["exit"] = "all rows empty after dropna"
            return ([], debug_info) if _debug else []

        # ── 10. Serialise ────────────────────────────────────
        records = []
        for _, row in data_df.iterrows():
            obj = {col: self._infer_value(row[col]) for col in data_df.columns}
            records.append(obj)

        debug_info["final_record_count"] = len(records)

        return (records, debug_info) if _debug else records

    # ── HELPERS ────────────────────────────────────────────────────────────

    @staticmethod
    def _is_empty(v):
        if pd.isna(v):
            return True
        s = str(v).strip()
        if s == "" or s.lower() == "nan":
            return True
        return False

    @staticmethod
    def _normalize(v):
        if IngestionEngine._is_empty(v):
            return ""
        v = str(v).replace("\n", " ").replace("\r", " ")
        v = re.sub(r"\s+", " ", v)
        return v.strip()

    @staticmethod
    def _safe_header(v):
        v = IngestionEngine._normalize(v)
        if v == "":
            return None
        v = v.lower()
        v = re.sub(r"[\/\\\-]+", "_", v)
        v = re.sub(r"[^a-zA-Z0-9_ ]", "", v)
        v = re.sub(r"\s+", "_", v)
        v = re.sub(r"_+", "_", v)
        return v.strip("_")

    @staticmethod
    def _infer_value(v):
        if IngestionEngine._is_empty(v):
            return None
        if isinstance(v, pd.Timestamp):
            return v.isoformat()
        if isinstance(v, (np.integer, int)):
            return int(v)
        if isinstance(v, (np.floating, float)):
            if pd.isna(v):
                return None
            if float(v).is_integer():
                return int(v)
            return float(v)
        
        s = str(v).strip()
        try:
            return int(s)
        except:
            pass
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except:
            pass
        try:
            dt = parse(s)
            if dt.year > 1900:
                return dt.isoformat()
        except:
            pass
        return s

    @staticmethod
    def _is_number_like(v):
        if IngestionEngine._is_empty(v):
            return False
        try:
            float(str(v))
            return True
        except:
            return False

    @staticmethod
    def _is_date_like(v):
        if IngestionEngine._is_empty(v):
            return False
        try:
            parse(str(v))
            return True
        except:
            return False

    # ── ENGINE COMPONENTS ──────────────────────────────────────────────────

    def _read_file(self, file):
        filename = getattr(file, "filename", "").lower()
        
        def _seek_back():
            if hasattr(file, "seek"):
                file.seek(0)

        try:
            if filename.endswith((".xlsx", ".xlsm")):
                df = pd.read_excel(file, header=None, engine="openpyxl", dtype=object)
                _seek_back()
                return df
            elif filename.endswith(".xls"):
                df = pd.read_excel(file, header=None, engine="xlrd", dtype=object)
                _seek_back()
                return df
            elif filename.endswith(".csv"):
                return pd.read_csv(file, header=None, dtype=object)
            else:
                df = pd.read_excel(file, header=None, engine="openpyxl", dtype=object)
                _seek_back()
                return df
        except Exception:
            _seek_back()
            return pd.read_csv(file, header=None, dtype=object)

    def _trim_dataframe(self, df):
        df = df.dropna(how="all")
        df = df.dropna(axis=1, how="all")
        return df.reset_index(drop=True)

    def _expand_header_merged_cells(self, df, header_rows_indices, file_obj=None):
        temp = df.copy()
        if file_obj is not None:
            fname = getattr(file_obj, "filename", "") or ""
            if fname.lower().endswith((".xlsx", ".xlsm")):
                try:
                    from openpyxl import load_workbook
                    import io as _io

                    if hasattr(file_obj, "seek"):
                        file_obj.seek(0)
                    raw = file_obj.read()
                    if hasattr(file_obj, "seek"):
                        file_obj.seek(0)

                    wb = load_workbook(_io.BytesIO(raw), data_only=True, read_only=False)
                    ws = wb.active

                    fill_map = {}
                    for merge_range in ws.merged_cells.ranges:
                        anchor_val = ws.cell(merge_range.min_row, merge_range.min_col).value
                        if anchor_val is None:
                            continue
                        for r in range(merge_range.min_row, merge_range.max_row + 1):
                            for c in range(merge_range.min_col, merge_range.max_col + 1):
                                fill_map[(r - 1, c - 1)] = anchor_val

                    header_set = set(header_rows_indices)
                    for (r, c), val in fill_map.items():
                        if r in header_set and r < len(temp) and c < temp.shape[1]:
                            temp.iloc[r, c] = val
                    return temp
                except Exception:
                    pass

        # Heuristic fallback
        n_cols = temp.shape[1]
        for i in header_rows_indices:
            if i >= len(temp): break
            row = temp.iloc[i].copy()
            last_val = None
            for j in range(n_cols):
                v = row.iloc[j]
                if not self._is_empty(v):
                    last_val = v
                elif last_val is not None:
                    if any(not self._is_empty(row.iloc[k]) for k in range(j + 1, n_cols)):
                        row.iloc[j] = last_val
            temp.iloc[i] = row
        return temp

    def _detect_header_rows(self, df):
        header_rows = []
        for i in range(min(self.max_header_scan, len(df))):
            row = df.iloc[i]
            vals = [row.iloc[j] for j in range(len(row)) if not self._is_empty(row.iloc[j])]
            if not vals: continue
            
            has_date = False
            for v in vals:
                if isinstance(v, (datetime.datetime, datetime.date)):
                    has_date = True; break
                s = str(v).strip()
                if (len(s) >= 6 and any(c.isdigit() for c in s) and any(c in s for c in "-/.")):
                    try:
                        parse(s); has_date = True; break
                    except: pass
            
            has_number = any(self._is_number_like(v) for v in vals)
            if has_date or has_number: break
            header_rows.append(i)
        return header_rows if header_rows else [0]

    def _build_headers(self, df, header_rows):
        header_df = df.iloc[header_rows].copy()
        headers = []
        for col_idx in range(df.shape[1]):
            parts = []
            for row_idx in range(len(header_rows)):
                cleaned = self._safe_header(header_df.iloc[row_idx, col_idx])
                if cleaned: parts.append(cleaned)
            
            final_parts = []
            for p in parts:
                if p not in final_parts: final_parts.append(p)

            if not final_parts:
                header_name = f"column_{col_idx}"
            elif len(final_parts) == 1:
                header_name = final_parts[0]
            else:
                parent, child = final_parts[0], final_parts[-1]
                header_name = f"{parent}_{child}" if child in ["g", "y", "r"] else "_".join(final_parts)
            headers.append(header_name)

        counts = {}
        final_headers = []
        for h in headers:
            if h not in counts:
                counts[h] = 1; final_headers.append(h)
            else:
                counts[h] += 1; final_headers.append(f"{h}_{counts[h]}")
        return final_headers

    def _detect_data_start(self, df, header_rows):
        start = max(header_rows) + 1
        for i in range(start, len(df)):
            if sum(1 for v in df.iloc[i] if not self._is_empty(v)) >= 2:
                return i
        return start

    def _should_fill_column(self, col_name, values):
        non_empty = [str(v).strip() for v in values if not self._is_empty(v)]
        if not non_empty: return False
        
        # 1. Pure text
        if any(self._is_number_like(v) or self._is_date_like(v) for v in non_empty):
            return False
        
        # 2. Sparse
        if (sum(1 for v in values if self._is_empty(v)) / len(values)) < 0.20:
            return False
        
        # 3. Repetitive
        if (len(set(non_empty)) / len(non_empty)) > 0.50:
            return False
        
        # 4. Short values
        if (sum(len(v) for v in non_empty) / len(non_empty)) > 60:
            return False
        
        # 5. Small vocabulary
        if len(set(non_empty)) > 20:
            return False
        
        return True

    def _fill_merged_data(self, data_df):
        df = data_df.copy()
        for col in df.columns:
            values = df[col].tolist()
            if not self._should_fill_column(col, values): continue
            filled, last_val = [], None
            for v in values:
                if not self._is_empty(v):
                    last_val = v; filled.append(v)
                else:
                    filled.append(last_val)
            df[col] = filled
        return df
